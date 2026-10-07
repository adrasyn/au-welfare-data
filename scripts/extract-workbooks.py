"""Read official workbooks and emit data for the static-site preparation step."""
from pathlib import Path
from collections import defaultdict
import json
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent
cache = root / '.cache' / 'sources'
population = {}
wb = load_workbook(cache / 'abs-population.xlsx', read_only=True, data_only=True)
for geo, sheet in [('lga', 'Table 1'), ('ced', 'Table 4')]:
    entries = {}
    for row in wb[sheet].iter_rows(min_row=8, values_only=True):
        code, name, value = row[0], row[1], row[25]
        if code is not None and str(code).isdigit() and isinstance(value, (int, float)):
            entries[str(code)] = {'name': name, 'population': int(value), 'populationDate': '2024-06-30', 'boundaryYear': 2024}
    population[geo] = entries
wb.close()

def rows(name):
    book = load_workbook(cache / f'{name}.xlsx', read_only=True, data_only=True)
    iterator = book.worksheets[0].iter_rows(values_only=True)
    headers = next(iterator)
    for values in iterator:
        yield dict(zip(headers, values))
    book.close()

poa_by_mb = {str(r['MB_CODE_2021']):str(r['POA_CODE_2021']).zfill(4) for r in rows('POA_2021_AUST')}
postcodes = defaultdict(lambda: {'ced':set(),'lga':set()})
registry = {'ced':{}, 'lga':{}}
for kind in ['CED','LGA']:
    geo = kind.lower()
    for r in rows(f'{kind}_2024_AUST'):
        code = str(r[f'{kind}_CODE_2024'])
        if code not in population[geo]:
            continue
        registry[geo][code] = {'name':r[f'{kind}_NAME_2024'], 'stateCode':str(r['STATE_CODE_2021']), 'stateName':r['STATE_NAME_2021']}
        pc = poa_by_mb.get(str(r['MB_CODE_2021']))
        if pc and pc.isdigit() and pc not in ['9494','9797']:
            postcodes[pc][geo].add(code)
    print(f'{geo}: {len(registry[geo])} regions')

out = {'population':population,'registry':registry,'postcodes':{pc:{g:sorted(ids) for g,ids in groups.items()} for pc,groups in sorted(postcodes.items())}}
(cache/'abs-extracted.json').write_text(json.dumps(out,separators=(',',':')))

book = load_workbook(cache/'dss-demographic.xlsx',read_only=True,data_only=True)
print('DSS sheets:',book.sheetnames)
for sheet in book.worksheets:
    if 'state' in sheet.title.lower() or 'gloss' in sheet.title.lower() or 'descr' in sheet.title.lower():
        content = [list(row) for row in sheet.iter_rows(values_only=True)]
        (cache/f'dss-{sheet.title.replace("/","-")}.json').write_text(json.dumps(content,default=str))
        print(sheet.title, len(content))
book.close()
